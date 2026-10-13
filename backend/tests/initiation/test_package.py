import csv
import io
from io import BytesIO

from docx import Document

from app.engagements.models import ClientContact
from app.files.models import FileObject
from app.files.service import read
from app.initiation.models import InitiationPackage, PackageDocument
from app.invitations.models import Invitation
from app.mail.models import OutboxMessage
from app.signoffs.models import SignOff
from app.tenancy.models import Organization, User
from app.tenancy.service import roles_of
from tests.factories import login_as
from tests.initiation.conftest import APPROVAL


def _package(client, cid):
    r = client.get(f"/api/v1/clients/{cid}/package")
    assert r.status_code == 200, r.text
    return r.json()


def test_package_has_7_documents_with_actions(client, db, people, submitted):
    cid, e = submitted()
    login_as(client, db, people["chief"])
    pkg = _package(client, cid)
    assert pkg["status"] == "draft" and pkg["version"] == 1
    assert [(d["item"], d["requires_action"]) for d in pkg["documents"]] == [
        (1, "sign"), (2, "nominate"), (3, "upload"), (4, "upload"), (5, "review"), (6, "read"), (7, "answer"),
    ]
    assert pkg["sends_to"] == [{"name": "Maria Lopez", "roles": ["Chief AI Compliance Officer", "Client Admin"], "email": "maria@abc.test"}]
    # opening it again does not make another package
    assert _package(client, cid)["id"] == pkg["id"]
    assert db.query(InitiationPackage).count() == 1


def test_package_docs_render_docx_and_pdf(client, db, people, submitted, settings):
    cid, e = submitted()
    login_as(client, db, people["chief"])
    pkg = _package(client, cid)
    letter = pkg["documents"][0]
    docx = read(settings, db.get(FileObject, letter["docx_file_id"]))
    text = "\n".join(p.text for p in Document(BytesIO(docx)).paragraphs)
    assert e.audit_ref in text and "Dear Maria Lopez" in text
    tables = Document(BytesIO(docx)).tables
    assert any("NYC Local Law 144" in c.text for t in tables for row in t.rows for c in row.cells)
    assert read(settings, db.get(FileObject, letter["pdf_file_id"])).startswith(b"%PDF")
    dept = pkg["documents"][3]
    rows = list(csv.reader(io.StringIO(read(settings, db.get(FileObject, dept["csv_file_id"])).decode())))
    assert rows[0] == ["department", "head_name", "head_email", "headcount"] and [r[0] for r in rows[1:]] == ["HR and recruiting", "Engineering and IT"]
    notice = client.get(f"/api/v1/packages/documents/{pkg['documents'][4]['id']}/preview").json()
    assert "52-c" in notice["html"]  # NY site in scope
    # files belong to the client tenant and the engagement
    f = db.get(FileObject, letter["pdf_file_id"])
    assert f.organization_id == cid and f.engagement_id == e.id and f.linked_type == "package_document"


def test_chief_can_edit_a_document_before_sending(client, db, people, submitted, settings):
    cid, _ = submitted()
    login_as(client, db, people["chief"])
    doc = _package(client, cid)["documents"][1]
    r = client.put(f"/api/v1/packages/documents/{doc['id']}", json={"body_md": "# Roles\n\nPlease name your Executive Sponsor by Friday."})
    assert r.status_code == 200
    new = r.json()
    assert new["edited"] is True and new["pdf_file_id"] != doc["pdf_file_id"]
    assert "Executive Sponsor by Friday" in client.get(f"/api/v1/packages/documents/{doc['id']}/preview").json()["html"]
    login_as(client, db, people["aud2"])
    assert client.put(f"/api/v1/packages/documents/{doc['id']}", json={"body_md": "x"}).status_code == 404


def test_approve_and_send_invites_caico_and_client_admin(client, db, people, submitted):
    cid, e = submitted(admin_same=False)
    login_as(client, db, people["chief"])
    _package(client, cid)
    r = client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL)
    assert r.status_code == 200, r.text
    db.expire_all()
    assert e.status == "initiated" and e.stage == 1 and e.approved_by == people["chief"].id
    assert db.get(Organization, cid).status == "active"
    maria = db.query(User).filter_by(email="maria@abc.test").one()
    sam = db.query(User).filter_by(email="sam@abc.test").one()
    assert {r.role for r in roles_of(db, maria.id)} == {"caico"} and {r.role for r in roles_of(db, sam.id)} == {"client_admin"}
    assert maria.organization_id == cid and maria.status == "invited"
    assert db.query(Invitation).filter(Invitation.user_id.in_([maria.id, sam.id])).count() == 2
    mail = db.query(OutboxMessage).filter(OutboxMessage.to == ["maria@abc.test"]).one()
    assert "We look forward to working with you." in mail.body_text and "/invite/" in mail.body_text
    assert db.query(ClientContact).filter_by(engagement_id=e.id, contact_role="caico").one().user_id == maria.id
    pkg = db.query(InitiationPackage).one()
    assert pkg.status == "sent" and pkg.approved_by == people["chief"].id and pkg.sent_at is not None
    so = db.query(SignOff).filter_by(engagement_id=e.id, number=1).one()
    assert so.status == "pending" and so.required_approver_roles == ["caico", "chief_auditor"] and so.preparer_id == people["aud"].id
    assert so.document_version_id == db.query(PackageDocument).filter_by(package_id=pkg.id, item=1).one().id


def test_same_person_caico_admin_single_invite(client, db, people, submitted):
    cid, _ = submitted(admin_same=True)
    login_as(client, db, people["chief"])
    _package(client, cid)
    assert client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL).status_code == 200
    maria = db.query(User).filter_by(email="maria@abc.test").one()
    assert {r.role for r in roles_of(db, maria.id)} == {"caico", "client_admin"}
    assert db.query(Invitation).filter_by(user_id=maria.id).count() == 1
    assert db.query(OutboxMessage).filter(OutboxMessage.to == ["maria@abc.test"]).count() == 1


def test_preparer_cannot_approve_package(client, db, people, submitted):
    from app.engagements.models import EngagementTeam

    cid, e = submitted()
    # the Chief Auditor on the team prepared it: another Chief Auditor must approve
    e.created_by = people["chief"].id
    db.add(EngagementTeam(organization_id=cid, engagement_id=e.id, user_id=people["chief2"].id, role="qa_reviewer"))
    db.commit()
    login_as(client, db, people["chief"])
    _package(client, cid)
    assert client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL).status_code == 403
    login_as(client, db, people["aud"])
    assert client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL).status_code == 403
    login_as(client, db, people["chief2"])
    assert client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL).status_code == 200


def test_approval_checklist_required(client, db, people, submitted):
    cid, e = submitted()
    login_as(client, db, people["chief"])
    _package(client, cid)
    r = client.post(f"/api/v1/clients/{cid}/approve", json={"client_message": "", "checklist": {"scope_ok": True, "no_conflict": True, "letter_previewed": False}})
    assert r.status_code == 422 and r.json()["errors"][0]["field"] == "checklist.letter_previewed"
    db.expire_all()
    assert e.status == "pending_approval"


def test_dashboard_auditor_sees_assigned_only(client, db, people, submitted):
    cid, e = submitted()
    login_as(client, db, people["aud2"])
    assert client.get("/api/v1/engagements").json()["rows"] == []
    login_as(client, db, people["aud"])
    d = client.get("/api/v1/engagements").json()
    assert [r["client_name"] for r in d["rows"]] == ["ABC Corp"]
    row = d["rows"][0]
    assert row["status_label"] == "Awaiting approval" and row["chief_auditor"] == "Cara Chief" and row["stage_label"] == "Draft"
    login_as(client, db, people["chief"])
    d = client.get("/api/v1/engagements").json()
    assert d["kpis"] == {"in_progress": 0, "completed": 0, "awaiting_approval": 1, "awaiting_client": 0}
    assert any("submitted" in a["summary"].lower() or "awaiting approval" in a["summary"].lower() for a in client.get("/api/v1/engagements/activity").json())
    _package(client, cid)
    client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL)
    d = client.get("/api/v1/engagements").json()
    assert d["kpis"]["awaiting_client"] == 1 and d["rows"][0]["status_label"] == "Initiation sent" and d["rows"][0]["stage_label"] == "1 · Onboarding"


def test_client_page_shows_whats_left_and_resend(client, db, people, submitted):
    cid, e = submitted()
    login_as(client, db, people["chief"])
    _package(client, cid)
    client.post(f"/api/v1/clients/{cid}/approve", json=APPROVAL)
    page = client.get(f"/api/v1/clients/{cid}").json()
    assert page["gate"]["unmet"] == ["Sign-off 1 (engagement letter and scope) is not approved yet."]
    caico = next(c for c in page["contacts"] if c["role"] == "caico")
    assert caico["invite_status"] == "Invitation sent, not yet accepted"
    assert any("NYC Local Law 144" in law["label"] for law in page["laws"])
    r = client.post(f"/api/v1/clients/{cid}/contacts/caico/resend")
    assert r.status_code == 200
    assert db.query(OutboxMessage).filter(OutboxMessage.to == ["maria@abc.test"]).count() == 2
    login_as(client, db, people["aud2"])
    assert client.get(f"/api/v1/clients/{cid}").status_code == 404


def test_client_cannot_see_engagement_before_it_is_sent(client, db, people, submitted):
    from tests.factories import make_user

    cid, e = submitted()
    early = make_user(db, db.get(Organization, cid), roles=["client_admin"], name="Early Bird", email="early@abc.test")
    db.commit()
    login_as(client, db, early)
    assert client.get("/api/v1/me/initiation").status_code == 404
    assert client.get(f"/api/v1/clients/{cid}/wizard").status_code == 404

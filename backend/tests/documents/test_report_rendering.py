"""Generated documents: fonts that draw every character, the report cover, contents, rating colours and score bars."""
import re
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.documents.docx import to_docx
from app.documents.pdf import FONTS, report_html, to_pdf
from app.documents.render import render_markdown

SYMBOLS = "– · → ≥ € ₹ ’ “quoted”"
META = {"audit_ref": "FSL-AI-AUD-2026-001"}
REPORT = f"""# AI governance audit report

**ABC Corp**

| Item | Detail |
|---|---|
| Audit reference | FSL-AI-AUD-2026-001 |

<!-- pagebreak -->

<!-- toc -->

<!-- pagebreak -->

## 1. Executive summary

Penalties up to €35M; INR 250 crore (₹250 crore); residual ≥ 70 → High; the client’s “final” view – A · B.

| Ref | Severity | Maturity (1–5) | Residual risk (0–100) |
|---|---|---|---|
| F-001 | High (4) | 2.25 | 58.2 |
| F-002 | Observation (1) | 3 | 10 |

## 2. Findings

Text.
"""


def _pdf_text(pdf: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)


def _fonts(pdf: bytes) -> set[str]:
    """The fonts that actually draw text (xhtml2pdf also lists Helvetica as an unused resource)."""
    out = set()

    def visit(text, cm, tm, font, size):
        if font and text.strip():
            out.add(str(font.get("/BaseFont")))

    for page in PdfReader(BytesIO(pdf)).pages:
        page.extract_text(visitor_text=visit)
    return out


def test_bundled_fonts_cover_report_symbols():
    for name in ("IBMPlexSans-Regular.ttf", "IBMPlexSans-Bold.ttf", "IBMPlexSans-Italic.ttf", "IBMPlexSans-BoldItalic.ttf", "DejaVuSans.ttf"):
        assert (FONTS / name).is_file(), name


def test_pdf_symbols_come_back_intact():
    pdf = to_pdf(f"# Symbols\n\n{SYMBOLS}\n\n**Bold {SYMBOLS}** and *italic {SYMBOLS}*\n", "Symbols", META)
    text = " ".join(_pdf_text(pdf).split())
    for ch in "–·→≥€₹’“”":
        assert ch in text, ch
    assert "■" not in text
    assert any("IBMPlexSans" in f for f in _fonts(pdf)), _fonts(pdf)
    assert all("IBMPlexSans" in f or "DejaVu" in f for f in _fonts(pdf)), _fonts(pdf)


def test_docx_symbols_come_back_intact():
    doc = Document(BytesIO(to_docx(f"# Symbols\n\n{SYMBOLS}\n", "Symbols", META)))
    assert SYMBOLS in "\n".join(p.text for p in doc.paragraphs)


def test_markers_never_show_as_text():
    html = render_markdown(REPORT)
    assert "<!--" not in html and "&lt;!--" not in html and "toc" not in re.sub(r"<[^>]+>", "", html).lower()
    text = _pdf_text(to_pdf(REPORT, "AI governance audit report v1.1", META | {"cover": True}))
    assert "pagebreak" not in text and "<!--" not in text and "**" not in text


def test_report_cover_contents_and_page_count():
    pdf = to_pdf(REPORT, "AI governance audit report v1.1", META | {"cover": True})
    reader = PdfReader(BytesIO(pdf))
    cover = reader.pages[0]
    assert cover.images, "the cover carries the Futran Solutions logo"
    assert "AI governance audit report" in cover.extract_text()
    contents = reader.pages[1].extract_text()
    assert "Contents" in contents and "Executive summary" in contents and "Findings" in contents
    assert re.search(r"Page \d+ of \d+", reader.pages[2].extract_text())


def test_rating_cells_are_coloured_and_scores_have_bars():
    html = report_html(REPORT, "AI governance audit report v1.1", META | {"cover": True})
    assert re.search(r'<td class="rating rating-high"[^>]*>\s*High \(4\)', html)
    assert re.search(r'<td class="rating rating-observation"[^>]*>\s*Observation \(1\)', html)
    assert html.count('class="bar-fill"') == 4  # maturity (1–5) and residual (0–100) cells
    assert 'class="bar-fill" style="width: 45%' in html  # 2.25 of 5
    assert 'style="width: 58%' in html  # 58.2 of 100


def test_docx_cover_logo_contents_and_rating_shading():
    doc = Document(BytesIO(to_docx(REPORT, "AI governance audit report v1.1", META | {"cover": True})))
    xml = doc.element.xml
    assert "<pic:pic" in xml, "the cover carries the logo"
    assert 'TOC \\o "2-2"' in xml
    cell = next(c for t in doc.tables for r in t.rows for c in r.cells if c.text == "High (4)")
    assert 'w:fill="B93A0B"' in cell._tc.xml
    hdr = doc.sections[0].header._element.xml
    assert "<pic:pic" in hdr or "Futran Solutions" in hdr

from io import BytesIO

from docx import Document

from app.documents.docx import to_docx
from app.documents.pdf import to_pdf
from app.documents.render import render_markdown

MD = """# Engagement letter

Dear **Maria**, this letter sets out the *scope* of the audit.

| System | Tier |
|---|---|
| HireScore | 3 |

- Engagement letter
- Roles note

<!-- pagebreak -->

## Scope

1. First item
2. Second item
"""
META = {"audit_ref": "FSL-AI-AUD-2026-001", "client": "ABC Corp"}


def test_render_markdown_tables_and_no_raw_html():
    html = render_markdown(MD + "\n<script>alert(1)</script>\n")
    assert "<table>" in html and "<h1>Engagement letter</h1>" in html
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert 'class="pagebreak"' in html


def test_docx_has_heading_table_and_footer_ref():
    doc = Document(BytesIO(to_docx(MD, "Engagement letter", META)))
    paras = doc.paragraphs
    h1 = next(p for p in paras if p.text == "Engagement letter")
    assert h1.style.name == "Heading 1"
    assert any(p.style.name == "Heading 2" and p.text == "Scope" for p in paras)
    dear = next(p for p in paras if p.text.startswith("Dear"))
    assert any(r.bold and r.text == "Maria" for r in dear.runs) and any(r.italic and r.text == "scope" for r in dear.runs)
    assert [p.text for p in paras if p.style.name == "List Bullet"] == ["Engagement letter", "Roles note"]
    assert [p.text for p in paras if p.style.name == "List Number"] == ["First item", "Second item"]
    t = doc.tables[0]
    assert t.cell(0, 0).text == "System" and t.cell(1, 1).text == "3"
    assert 'w:type="page"' in doc.element.xml
    sec = doc.sections[0]
    assert "Confidential" in sec.header.paragraphs[0].text and "<pic:pic" in sec.header._element.xml  # the Futran Solutions logo
    footer_xml = sec.footer._element.xml
    assert "FSL-AI-AUD-2026-001" in footer_xml and "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_pdf_starts_with_pdf_and_has_pages():
    pdf = to_pdf(MD, "Engagement letter", META)
    assert pdf.startswith(b"%PDF")
    assert pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages") >= 2  # the page break gives a second page


def test_unicode_survives_both_formats():
    md = "# Notice\n\nNY Civil Rights Law §52-c — employee monitoring. Café, Zürich, “quotes”.\n"
    doc = Document(BytesIO(to_docx(md, "Notice", META)))
    assert "§52-c — employee" in doc.paragraphs[1].text
    assert to_pdf(md, "Notice", META).startswith(b"%PDF")

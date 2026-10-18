"""Task 30: the audit report (template §12–13, spec §18.7–18.8), declarations, sign-offs 7 and 8 and the public summary."""
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import select

from app.engagements.models import ClientRequirement, Engagement, EngagementTeam
from app.laws.models import Regulation
from app.signoffs.service import current
from tests.factories import approve_client_review, login_as, make_user
from tests.initiation.conftest import live, people, submitted  # noqa: F401
from tests.inventory.test_named_users import inv  # noqa: F401

NOT_AGREEMENT = "Signing confirms you have received and read this report. It does not mean you agree with every finding."
SUMMARY = {"most_significant_text": "Recruiters paste applicant data into personal ChatGPT accounts.", "strengths": "A named AI owner and an approved AI use policy.",
           "immediate_actions": "Block personal AI accounts on company devices.", "narrative": "Governance is defined but not yet measured."}


@pytest.fixture
def ready(client, db, inv, people):
    """One approved finding, the engagement at stage 6, sign-off 5 not yet sent."""
    cid, e, maria, *_ = inv
    db.get(Engagement, e.id).stage = 5
    db.commit()
    approve_client_review(db, e)
    login_as(client, db, people["aud"])
    f = client.post(f"/api/v1/clients/{cid}/findings", json={"title": "Personal AI accounts used with applicant data", "severity": "high",
                                                             "description": "Recruiters use personal accounts.", "recommended_remediation": "Move to company accounts.", "proposed_owner": "Head of HR"}).json()
    client.post(f"/api/v1/findings/{f['id']}/submit")
    login_as(client, db, people["chief"])
    assert client.post(f"/api/v1/findings/{f['id']}/approve", json={}).json()["status"] == "approved"
    db.get(Engagement, e.id).stage = 6
    db.commit()
    return cid, e, maria, f


def _so5(client, db, people, cid):
    login_as(client, db, people["aud"])
    so = client.post(f"/api/v1/clients/{cid}/signoffs/5/submit", json={}).json()
    login_as(client, db, people["chief"])
    assert client.post(f"/api/v1/signoffs/{so['id']}/sign", json={"typed_name": people["chief"].name}).json()["status"] == "approved"


def _no_public_summary(db, e):
    """The fixture client is in New York, so NYC LL144 is in scope; take every law out of scope."""
    for r in db.scalars(select(ClientRequirement).where(ClientRequirement.engagement_id == e.id)):
        r.in_scope = False
    db.commit()


def _generate(client, cid, version):
    return client.post(f"/api/v1/clients/{cid}/report/generate", json={"version": version})


def _declare(client, db, user, report_id, kind):
    login_as(client, db, user)
    return client.post(f"/api/v1/reports/{report_id}/declarations", json={"kind": kind, "typed_name": user.name})


def _released(client, db, people, cid):
    """Sign-off 5, the lead review version, both Futran declarations and the v1.0 release."""
    _so5(client, db, people, cid)
    login_as(client, db, people["aud"])
    assert client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY).status_code == 200
    r = _generate(client, cid, "v0.2")
    assert r.status_code == 200, r.text
    rid = r.json()["id"]
    assert _declare(client, db, people["chief"], rid, "lead_auditor").status_code == 200
    assert _declare(client, db, people["em"], rid, "engagement_manager").status_code == 200
    login_as(client, db, people["em"])
    out = client.post(f"/api/v1/clients/{cid}/report/release")
    assert out.status_code == 200, out.text
    return out.json()


def test_report_v1_0_requires_declarations_and_so5(client, db, ready, people):
    cid, e, maria, f = ready
    login_as(client, db, people["aud"])
    client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY)
    r = _generate(client, cid, "v0.2")
    assert r.status_code == 409 and "Sign-off 5" in r.text
    login_as(client, db, people["em"])
    r = client.post(f"/api/v1/clients/{cid}/report/release")
    assert r.status_code == 409 and any("Sign-off 5" in m["message"] for m in r.json()["errors"])

    _so5(client, db, people, cid)
    login_as(client, db, people["aud"])
    rid = _generate(client, cid, "v0.2").json()["id"]
    login_as(client, db, people["em"])
    r = client.post(f"/api/v1/clients/{cid}/report/release")
    msgs = " ".join(m["message"] for m in r.json()["errors"])
    assert r.status_code == 409 and "Lead Auditor" in msgs and "Engagement Manager" in msgs
    assert _declare(client, db, people["aud"], rid, "lead_auditor").status_code == 403  # only the engagement's Lead Auditor
    assert _declare(client, db, people["chief"], rid, "lead_auditor").status_code == 200
    assert _declare(client, db, people["chief"], rid, "lead_auditor").status_code == 409  # once
    assert _declare(client, db, people["em"], rid, "engagement_manager").status_code == 200
    login_as(client, db, maria)
    assert client.get(f"/api/v1/clients/{cid}/report").json()["reports"] == []  # nothing released yet
    login_as(client, db, people["em"])
    out = client.post(f"/api/v1/clients/{cid}/report/release")
    assert out.status_code == 200, out.text
    assert out.json()["version"] == "v1.0" and out.json()["released_to_client_at"]
    db.expire_all()
    assert db.get(Engagement, e.id).stage == 7  # the draft report goes to the client: stage 6 is done
    login_as(client, db, maria)
    seen = client.get(f"/api/v1/clients/{cid}/report").json()["reports"]
    assert [x["version"] for x in seen] == ["v1.0"]
    assert client.get(f"/api/v1/files/{seen[0]['pdf_file_id']}").status_code == 200
    login_as(client, db, people["aud"])
    lead = client.get(f"/api/v1/clients/{cid}/report").json()
    v02 = next(x for x in lead["reports"] if x["version"] == "v0.2")
    login_as(client, db, maria)
    assert client.get(f"/api/v1/files/{v02['docx_file_id']}").status_code == 404  # the lead review copy stays with Futran


def test_changed_report_needs_new_declarations(client, db, ready, people):
    cid, e, maria, f = ready
    _so5(client, db, people, cid)
    login_as(client, db, people["aud"])
    client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY)
    rid = _generate(client, cid, "v0.2").json()["id"]
    _declare(client, db, people["chief"], rid, "lead_auditor")
    _declare(client, db, people["em"], rid, "engagement_manager")
    login_as(client, db, people["aud"])
    client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY | {"strengths": "Something else."})
    login_as(client, db, people["em"])
    r = client.post(f"/api/v1/clients/{cid}/report/release")
    assert r.status_code == 409 and "changed since" in r.text


def test_qa_declaration_optional_without_qa_reviewer(client, db, ready, people):
    cid, e, maria, f = ready
    out = _released(client, db, people, cid)
    from app.reports.models import Report

    v10 = db.get(Report, out["id"])
    assert "No QA Reviewer assigned" in v10.body_md


def test_qa_declaration_required_with_qa_reviewer(client, db, ready, people):
    cid, e, maria, f = ready
    from app.tenancy.models import Organization

    q = make_user(db, db.get(Organization, people["chief"].organization_id), roles=["qa_reviewer"], name="Quinn Reviewer", email="quinn@futran.test")
    db.add(EngagementTeam(organization_id=cid, engagement_id=e.id, user_id=q.id, role="qa_reviewer"))
    db.commit()
    _so5(client, db, people, cid)
    login_as(client, db, people["aud"])
    client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY)
    rid = _generate(client, cid, "v0.2").json()["id"]
    _declare(client, db, people["chief"], rid, "lead_auditor")
    _declare(client, db, people["em"], rid, "engagement_manager")
    login_as(client, db, people["em"])
    r = client.post(f"/api/v1/clients/{cid}/report/release")
    assert r.status_code == 409 and "Quality Reviewer" in r.text
    assert _declare(client, db, q, rid, "quality_reviewer").status_code == 200
    login_as(client, db, people["em"])
    assert client.post(f"/api/v1/clients/{cid}/report/release").status_code == 200


def test_executive_summary_requires_strengths(client, db, ready, people):
    cid, e, maria, f = ready
    _so5(client, db, people, cid)
    login_as(client, db, people["aud"])
    assert client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY | {"strengths": "  "}).status_code == 200  # a draft may be incomplete
    assert _generate(client, cid, "v0.1").status_code == 200
    r = _generate(client, cid, "v0.2")
    assert r.status_code == 422 and r.json()["errors"][0]["field"] == "strengths"
    assert "genuine positives" in r.json()["errors"][0]["message"]


def test_report_docx_contains_findings_table_and_scorecard(client, db, ready, people):
    cid, e, maria, f = ready
    login_as(client, db, people["aud"])
    client.put(f"/api/v1/clients/{cid}/report/summary", json=SUMMARY)
    r = client.post(f"/api/v1/clients/{cid}/report/roadmap", json={"action": "Move recruiters to company AI accounts", "owner": "Head of HR", "priority": "high",
                                                                   "target_date": "2027-01-31", "how_futran_can_help": "AI acceptable use policy workshop", "finding_ids": [f["id"]]})
    assert r.status_code == 200 and r.json()["ref"] == "R-01"
    rep = _generate(client, cid, "v0.1").json()
    assert rep["version"] == "v0.1"
    data = client.get(f"/api/v1/files/{rep['docx_file_id']}").content
    doc = Document(BytesIO(data))
    headers = [[c.text for c in t.rows[0].cells] for t in doc.tables]
    assert ["Ref", "Finding", "Severity", "System", "Owner", "Due"] in headers
    assert any(h[:2] == ["Domain", "Maturity (1–5)"] for h in headers)
    assert any("How Futran can help" in h for h in headers)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "F-001" in text and "R-01" in text and "Move recruiters to company AI accounts" in text
    assert "Appendix B" in text and "Appendix C" in text and "Disparate Impact" in text
    assert "Greatest strengths observed" in text and "A named AI owner" in text
    assert client.get(f"/api/v1/files/{rep['pdf_file_id']}").content[:4] == b"%PDF"


def _final(client, db, people, cid, e):
    db.get(Engagement, e.id).stage = 8  # sign-off 6 (Task 31) is approved
    db.commit()
    login_as(client, db, people["aud"])
    r = client.post(f"/api/v1/clients/{cid}/report/final")
    assert r.status_code == 200, r.text
    so = r.json()["sign_off"]
    assert [p["role"] for p in so["parties"]] == ["chief_auditor"]
    login_as(client, db, people["chief"])
    assert client.post(f"/api/v1/signoffs/{so['id']}/sign", json={"typed_name": people["chief"].name}).json()["status"] == "approved"


def test_client_ack_text_states_not_agreement(client, db, ready, people):
    cid, e, maria, f = ready
    _no_public_summary(db, e)
    _released(client, db, people, cid)
    login_as(client, db, maria)
    assert client.post(f"/api/v1/clients/{cid}/report/acknowledge", json={"typed_name": "Maria Lopez", "commitment": "plan"}).status_code == 409  # not final yet
    _final(client, db, people, cid, e)
    login_as(client, db, maria)
    view = client.get(f"/api/v1/clients/{cid}/report").json()
    assert [x["version"] for x in view["reports"]] == ["v1.1", "v1.0"]
    ack = view["acknowledgement"]
    assert ack["notice"] == NOT_AGREEMENT and ack["can_sign"]
    assert "has received the Futran Solutions AI Governance Audit Report ref " + e.audit_ref in ack["text"]
    bad = client.post(f"/api/v1/clients/{cid}/report/acknowledge", json={"typed_name": "Maria Lopez", "commitment": "disputes", "disputed_refs": ["F-999"]})
    assert bad.status_code == 422
    r = client.post(f"/api/v1/clients/{cid}/report/acknowledge", json={"typed_name": "Maria Lopez", "commitment": "disputes", "disputed_refs": [f["ref"]]})
    assert r.status_code == 200, r.text
    assert r.json()["acknowledgement"]["signed_by"] == "Maria Lopez" and r.json()["acknowledgement"]["disputed_refs"] == ["F-001"]
    db.expire_all()
    assert db.get(Engagement, e.id).stage == 9  # no public summary needed: stage 8 is done


def test_public_summary_only_when_ll144(client, db, ready, people):
    cid, e, maria, f = ready
    login_as(client, db, people["aud"])
    assert client.get(f"/api/v1/clients/{cid}/report").json()["public_summary"]["required"] is True  # NYC LL144 is in scope
    _no_public_summary(db, e)
    assert client.get(f"/api/v1/clients/{cid}/report").json()["public_summary"]["required"] is False
    assert client.put(f"/api/v1/clients/{cid}/public-summary", json={"tool_name": "Résumé screener"}).status_code == 409
    reg = db.scalar(select(Regulation).where(Regulation.public_summary.is_(False)).limit(1))
    reg.public_summary = True  # an Admin flags a law as needing a public summary
    db.add(ClientRequirement(organization_id=cid, engagement_id=e.id, regulation_id=reg.id, source="manual", in_scope=True))
    db.commit()
    assert client.get(f"/api/v1/clients/{cid}/report").json()["public_summary"]["required"] is True
    _released(client, db, people, cid)
    _final(client, db, people, cid, e)
    login_as(client, db, maria)
    client.post(f"/api/v1/clients/{cid}/report/acknowledge", json={"typed_name": "Maria Lopez", "commitment": "plan"})
    db.expire_all()
    assert db.get(Engagement, e.id).stage == 8  # sign-off 8 is still needed
    from app.engagements.gates import gate_status

    assert any("Sign-off 8" in u for u in gate_status(db, db.get(Engagement, e.id)).unmet)
    login_as(client, db, people["aud"])
    assert client.post(f"/api/v1/clients/{cid}/signoffs/8/submit", json={}).status_code == 409  # no impact ratios yet
    body = {"tool_name": "Résumé screener", "audit_date": "2026-11-02", "distribution_date": "2026-11-20", "data_source": "Historical applicant data, Jan–Jun 2026",
            "impact_ratios": [{"category": "Sex", "group": "Female", "selection_rate": 0.42, "impact_ratio": 0.93}, {"category": "Sex", "group": "Male", "selection_rate": 0.45, "impact_ratio": 1.0}]}
    r = client.put(f"/api/v1/clients/{cid}/public-summary", json=body)
    assert r.status_code == 200, r.text
    bad = client.put(f"/api/v1/clients/{cid}/public-summary", json=body | {"impact_ratios": [{"category": "Sex", "group": "F", "selection_rate": 2, "impact_ratio": -1}]})
    assert bad.status_code == 422
    so = client.post(f"/api/v1/clients/{cid}/signoffs/8/submit", json={}).json()
    assert [p["role"] for p in so["parties"]] == ["caico"]
    login_as(client, db, maria)
    assert client.post(f"/api/v1/signoffs/{so['id']}/sign", json={"typed_name": "Maria Lopez"}).json()["status"] == "approved"
    db.expire_all()
    assert db.get(Engagement, e.id).stage == 9
    assert current(db, db.get(Engagement, e.id), 8).status == "approved"

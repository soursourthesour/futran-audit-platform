"""Parse the Futran Global AI Law Matrix workbook into a reviewable preview (pure: no database).

Sheets (header on row 4, data from row 5, a row is data when column A is filled):
  Federal Laws                  -> regulations under US
  State Matrix                  -> 26 US states x 9 features; each cell split on line breaks
  Global Matrix                 -> per-jurisdiction feature notes (summaries, not regulations)
  EU, UK, Canada, Mexico, Brazil, India, Australia, Singapore -> regulations, same columns as Federal Laws
  Gulf                          -> like the country sheets; "<Country> · <Category>" names the country
  Read Me                       -> status date ("Status as of 5 October 2026")
  Sources                       -> citations kept on the matrix version
"""
import hashlib
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pycountry
import yaml

SEEDS = Path(__file__).resolve().parents[2] / "seeds"

FEATURES = [
    "consequential_decisions_admt",
    "employment_hiring",
    "privacy_profiling",
    "health_insurance",
    "chatbots_minors",
    "synthetic_content",
    "frontier_safety",
    "biometrics_pricing",
    "government_other",
]
FEATURE_LABELS = {
    "consequential_decisions_admt": "Consequential decisions / ADMT",
    "employment_hiring": "Employment and hiring AI",
    "privacy_profiling": "Privacy: profiling and automated decisions",
    "health_insurance": "Health care and insurance AI",
    "chatbots_minors": "Chatbots, companions and minors",
    "synthetic_content": "Synthetic content, deepfakes and likeness",
    "frontier_safety": "Frontier model safety",
    "biometrics_pricing": "Biometrics, neural data and algorithmic pricing",
    "government_other": "Government use and other",
}
TYPES = ["statute", "regulation", "executive", "voluntary", "pending", "policy", "guidance"]

COUNTRY_SHEETS = {"EU": "EU", "UK": "GB", "Canada": "CA", "Mexico": "MX", "Brazil": "BR", "India": "IN", "Australia": "AU", "Singapore": "SG"}
GLOBAL_NAMES = {
    "European Union": "EU", "United Kingdom": "GB", "Canada": "CA", "Mexico": "MX", "Brazil": "BR", "India": "IN",
    "UAE": "AE", "United Arab Emirates": "AE", "Saudi Arabia": "SA", "Qatar": "QA", "Bahrain": "BH", "Oman": "OM",
    "Kuwait": "KW", "Australia": "AU", "Singapore": "SG",
}
GLOBAL_NAMES_DISPLAY = {v: k for k, v in GLOBAL_NAMES.items() if k not in ("UAE",)}

NOTE_PREFIXES = ("key date", "vetoed", "2026 session", "2025 session", "pending:", "watch:")
CONTINUATION_PREFIXES = ("developers:", "deployers:", "also ", "note:", "exemptions", "penalties", "enforcement:", "applies ", "covers ")
LAW_START = re.compile(r"^(?:[A-Z]{1,5}\.?\s?){1,2}\d|^[A-Z]{3,}\b")
LAW_WORDS = re.compile(r"\b(?:Act|Law|Laws|Rules?|Regulations?|Code|Order|Ordinance|Guidance|Bill|Policy|Directive|Statute)\b")
NYC_PATTERN = re.compile(r"local law 144|\bnyc\b|new york city", re.I)
DEVELOPER_PATTERN = re.compile(r"frontier|developers only|frontier developers|\bdevelopers? of (?:frontier|large)", re.I)


@dataclass
class PreviewRow:
    row_key: str
    kind: str  # regulation | note
    source_sheet: str
    jurisdiction_code: str
    jurisdiction_name: str
    level: str
    feature: str
    short_name: str
    full_text: str
    type: str = "statute"
    type_text: str | None = None
    category_text: str | None = None
    enforcer: str | None = None
    trigger_text: str | None = None
    obligations_text: str | None = None
    developments_text: str | None = None
    most_affected: str | None = None
    verify_flag: bool = False
    applies_to: str = "deployer"

    def content_hash(self) -> str:
        parts = [self.short_name, self.full_text, self.type, self.enforcer or "", self.trigger_text or "", self.obligations_text or "", self.developments_text or "", self.most_affected or ""]
        return hashlib.sha256("\x1f".join(parts).encode()).hexdigest()


@dataclass
class ImportPreview:
    status_date: date | None
    source_filename: str
    rows: list[PreviewRow] = field(default_factory=list)
    jurisdiction_notes: dict[str, dict[str, str]] = field(default_factory=dict)
    sources: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return {
            "status_date": self.status_date.isoformat() if self.status_date else None,
            "source_filename": self.source_filename,
            "rows": [asdict(r) for r in self.rows],
            "jurisdiction_notes": self.jurisdiction_notes,
            "sources": self.sources,
            "warnings": self.warnings,
        }

    @classmethod
    def from_json(cls, data: dict) -> "ImportPreview":
        return cls(
            status_date=date.fromisoformat(data["status_date"]) if data.get("status_date") else None,
            source_filename=data["source_filename"],
            rows=[PreviewRow(**r) for r in data["rows"]],
            jurisdiction_notes=data["jurisdiction_notes"],
            sources=data["sources"],
            warnings=data["warnings"],
        )


def _feature_rules() -> tuple[list[tuple[re.Pattern, str]], str]:
    data = yaml.safe_load((SEEDS / "feature_map.yaml").read_text())
    return [(re.compile(r["match"], re.I), r["feature"]) for r in data["rules"]], data["default"]


def map_feature(category: str, warnings: list[str] | None = None) -> str:
    rules, default = _feature_rules()
    for pat, feat in rules:
        if pat.search(category or ""):
            return feat
    if warnings is not None:
        warnings.append(f"Category '{category}' has no feature rule; filed under {FEATURE_LABELS[default]}.")
    return default


def map_type(text: str | None) -> str:
    t = (text or "").lower()
    if re.search(r"pending|draft|died|bill", t):
        return "pending"
    if "executive" in t:
        return "executive"
    if t.startswith("guidance") or t.startswith("supervisory") or t.startswith("regulator guidance") or t.startswith("regulator framework"):
        return "guidance"
    if re.search(r"voluntary|pledge|principles|standard", t) and "binding" not in t:
        return "voluntary"
    if t.startswith("policy") or "agency policy" in t or t.startswith("mandatory policy"):
        return "policy"
    if re.search(r"regulation|rules|directive|regulator", t) and "statute" not in t:
        return "regulation"
    return "statute"


def short_name_of(line: str) -> str:
    text = line.strip()
    cut = len(text)
    for sep in (" (", ":", " — ", "; "):
        i = text.find(sep)
        if 0 < i < cut:
            cut = i
    name = text[:cut].strip().rstrip(".")
    if len(name) > 120:
        name = name[:117].rsplit(" ", 1)[0] + "…"
    return name


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return None if s in ("", "—", "-") else s


def _state_code(name: str, warnings: list[str]) -> str | None:
    base = re.sub(r"\s*\(.*\)\s*$", "", name).strip()
    for s in pycountry.subdivisions.get(country_code="US"):
        if s.name.lower() == base.lower():
            return s.code
    warnings.append(f"State '{name}' not recognised.")
    return None


def _status_date(wb) -> date | None:
    for row in wb["Read Me"].iter_rows(values_only=True):
        for cell in row:
            if cell and "Status as of" in str(cell):
                m = re.search(r"Status as of (\d{1,2} \w+ \d{4})", str(cell))
                if m:
                    return datetime.strptime(m.group(1), "%d %B %Y").date()
    return None


def _data_rows(ws):
    for idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        if row and row[0] not in (None, ""):
            yield idx, row


def _law_row(sheet: str, idx: int, row, code: str, name: str, category: str, warnings) -> PreviewRow:
    law = _clean(row[1]) or ""
    texts = " ".join(str(c) for c in row[1:8] if c)
    feature = map_feature(category, warnings)
    applies = "developer" if feature == "frontier_safety" or DEVELOPER_PATTERN.search(" ".join(str(c) for c in (row[4], row[7]) if c)) else "deployer"
    return PreviewRow(
        row_key=f"{sheet}:{idx}",
        kind="regulation",
        source_sheet=sheet,
        jurisdiction_code=code,
        jurisdiction_name=name,
        level="supranational" if code == "EU" else "country",
        feature=feature,
        short_name=short_name_of(law) if len(law) > 120 else law,
        full_text=law,
        type=map_type(_clean(row[2])),
        type_text=_clean(row[2]),
        category_text=category,
        enforcer=_clean(row[3]),
        trigger_text=_clean(row[4]),
        obligations_text=_clean(row[5]),
        developments_text=_clean(row[6]),
        most_affected=_clean(row[7]),
        verify_flag="(verify)" in texts,
        applies_to=applies,
    )


def parse_workbook(path: str | Path) -> ImportPreview:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pv = ImportPreview(status_date=_status_date(wb), source_filename=path.name)
    w = pv.warnings

    # Federal Laws
    for idx, row in _data_rows(wb["Federal Laws"]):
        pv.rows.append(_law_row("Federal Laws", idx, row, "US", "United States", str(row[0]), w))

    # State Matrix
    for idx, row in _data_rows(wb["State Matrix"]):
        state_name = str(row[0])
        code = _state_code(state_name, w)
        if not code:
            continue
        clean_name = re.sub(r"\s*\(.*\)\s*$", "", state_name)
        for col, feature in enumerate(FEATURES, start=1):
            cell = _clean(row[col] if col < len(row) else None)
            if not cell:
                continue
            current: PreviewRow | None = None
            for li, raw in enumerate(cell.split("\n")):
                line = raw.strip()
                if not line:
                    continue
                low = line.lower()
                key = f"State Matrix:{idx}:{col}:{li}"
                if low.startswith(NOTE_PREFIXES) or (feature == "government_other" and not re.match(r"^(?:[A-Z]{1,4}\s?\d|[A-Z][A-Za-z.]+ (?:Act|Law|Rule|Order))", line)):
                    pv.rows.append(PreviewRow(key, "note", "State Matrix", code, clean_name, "state", feature, line[:120], line, verify_flag="(verify)" in line))
                    current = None
                    continue
                starts_law = bool(LAW_START.match(line) or LAW_WORDS.search(short_name_of(line)))
                if current is not None and (low.startswith(CONTINUATION_PREFIXES) or line[0].islower() or not starts_law):
                    current.full_text += "\n" + line
                    current.verify_flag = current.verify_flag or "(verify)" in line
                    continue
                if low.startswith("no ") and ". " in line:
                    first, rest = line.split(". ", 1)
                    pv.rows.append(PreviewRow(key + "n", "note", "State Matrix", code, clean_name, "state", feature, first[:120], first + "."))
                    line = rest
                jcode, jname, level = code, clean_name, "state"
                if code == "US-NY" and NYC_PATTERN.search(line):
                    jcode, jname, level = "US-NYC", "New York City", "city"
                current = PreviewRow(
                    row_key=key,
                    kind="regulation",
                    source_sheet="State Matrix",
                    jurisdiction_code=jcode,
                    jurisdiction_name=jname,
                    level=level,
                    feature=feature,
                    short_name=short_name_of(line),
                    full_text=line,
                    type="pending" if re.search(r"pending|sent to the governor|introduced|bill \(verify\)", low) else "statute",
                    category_text=FEATURE_LABELS[feature],
                    verify_flag="(verify)" in line,
                    applies_to="developer" if feature == "frontier_safety" or DEVELOPER_PATTERN.search(line) else "deployer",
                )
                pv.rows.append(current)

    # Global Matrix — notes per jurisdiction and feature
    for idx, row in _data_rows(wb["Global Matrix"]):
        name = str(row[0]).strip()
        code = GLOBAL_NAMES.get(name)
        if not code:
            w.append(f"Global Matrix jurisdiction '{name}' not recognised.")
            continue
        notes = {}
        for col, feature in enumerate(FEATURES, start=1):
            cell = _clean(row[col] if col < len(row) else None)
            if cell:
                notes[feature] = cell
        pv.jurisdiction_notes[code] = notes

    # Country sheets
    for sheet, code in COUNTRY_SHEETS.items():
        if sheet not in wb.sheetnames:
            w.append(f"Sheet '{sheet}' is missing.")
            continue
        name = GLOBAL_NAMES_DISPLAY.get(code, sheet)
        for idx, row in _data_rows(wb[sheet]):
            pv.rows.append(_law_row(sheet, idx, row, code, name, str(row[0]), w))

    # Gulf
    if "Gulf" in wb.sheetnames:
        for idx, row in _data_rows(wb["Gulf"]):
            cat = str(row[0])
            country, _, category = cat.partition(" · ")
            code = GLOBAL_NAMES.get(country.strip())
            if not code:
                w.append(f"Gulf row {idx}: country '{country}' not recognised.")
                continue
            pv.rows.append(_law_row("Gulf", idx, row, code, GLOBAL_NAMES_DISPLAY.get(code, country), category.strip() or cat, w))

    # Sources
    if "Sources" in wb.sheetnames:
        for _, row in _data_rows(wb["Sources"]):
            topic, src = str(row[0]), str(row[1] or "")
            m = re.search(r"https?://\S+", src)
            pv.sources.append({"topic": topic, "source": src.split(" — ")[0].strip(), "url": m.group(0) if m else None})
    wb.close()
    return pv

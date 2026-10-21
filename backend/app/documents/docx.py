"""Markdown → DOCX with python-docx: headings 1–3, paragraphs, bold/italic, lists, tables, quotes and page breaks.

Every document carries the Futran Solutions logo in its header and "Page n of m" in its footer. With meta["cover"], the
text before the first page break becomes a cover page (logo, no running header), `<!-- toc -->` becomes a Word table of
contents, and rating cells such as "High (4)" are shaded with the risk level colour.
"""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.documents.pdf import LOGO
from app.documents.render import SENTINEL, TOC_SENTINEL, bar_pct, bar_scale, column_widths, tokens
from app.scoring.scale import from_labelled

NAVY = RGBColor(0x0B, 0x2E, 0x6A)
GREY = RGBColor(0x5B, 0x64, 0x74)
FONT = "Arial"  # IBM Plex Sans is the brand face, but Arial is on every reader's machine
HEADER_FILL = "E6EDF7"
TEXT_WIDTH_CM = 17.1  # US Letter less 2.2 cm margins


def _el(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


def _field(paragraph, instr: str, size: Pt | None = None) -> None:
    run = paragraph.add_run()
    if size:
        run.font.size = size
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    it = _el("w:instrText", **{"xml:space": "preserve"})
    it.text = f" {instr} "
    run._r.append(it)
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    t = paragraph.add_run("1")
    if size:
        t.font.size = size
    end = paragraph.add_run()
    end._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def _inline(paragraph, inline_token) -> None:
    bold = italic = False
    for t in inline_token.children or []:
        if t.type == "strong_open":
            bold = True
        elif t.type == "strong_close":
            bold = False
        elif t.type == "em_open":
            italic = True
        elif t.type == "em_close":
            italic = False
        elif t.type in ("text", "code_inline"):
            r = paragraph.add_run(t.content)
            r.bold, r.italic = bold or None, italic or None
            if t.type == "code_inline":
                r.font.name = "Consolas"
        elif t.type == "softbreak":
            paragraph.add_run(" ")
        elif t.type == "hardbreak":
            paragraph.add_run().add_break()


def _styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Header", "Footer"):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(8), GREY
    for i, size in ((1, 20), (2, 14), (3, 11.5)):
        s = doc.styles[f"Heading {i}"]
        s.font.name, s.font.size, s.font.color.rgb, s.font.bold = FONT, Pt(size), NAVY, True
        fonts = s.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # the theme font would win over Arial
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(14 if i < 3 else 10)
        s.paragraph_format.space_after = Pt(5)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    tc_pr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(_el("w:tblHeader", **{"w:val": "true"}))


def _right_tab(p) -> None:
    """One right-aligned tab at the margin; the Header and Footer styles' own centre and right stops are cleared."""
    stops = p.paragraph_format.tab_stops
    stops.add_tab_stop(Cm(8.255), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Cm(16.51), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)


def _running(doc, title: str, meta: dict) -> None:
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    hp = sec.header.paragraphs[0]
    _right_tab(hp)
    if LOGO.is_file():
        hp.add_run().add_picture(str(LOGO), height=Cm(0.45))
    else:
        hp.add_run("Futran Solutions").font.size = Pt(8)
    conf = hp.add_run("\tConfidential")
    conf.font.size, conf.font.color.rgb = Pt(8), GREY
    fp = sec.footer.paragraphs[0]
    _right_tab(fp)
    ref = str(meta.get("audit_ref", "") or "")
    left = fp.add_run(f"{ref}   {title}" if ref else title)
    left.font.size, left.font.color.rgb = Pt(8), GREY
    fp.add_run("\tPage ").font.size = Pt(8)
    _field(fp, "PAGE", Pt(8))
    fp.add_run(" of ").font.size = Pt(8)
    _field(fp, "NUMPAGES", Pt(8))
    if meta.get("cover"):
        sec.different_first_page_header_footer = True


def _toc(doc, headings: list[str]) -> None:
    """A Word contents field. Word fills in the page numbers when the document is opened (it asks to update fields);
    until then the field shows the section headings."""
    title = doc.add_paragraph()
    r = title.add_run("Contents")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), NAVY
    first = doc.add_paragraph()
    run = first.add_run()
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    it = _el("w:instrText", **{"xml:space": "preserve"})
    it.text = ' TOC \\o "2-2" \\h \\z \\u '
    run._r.append(it)
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    paras = [first] + [doc.add_paragraph() for _ in headings[1:]]
    for p, h in zip(paras, headings or ["Update this field to show the contents."]):
        p.add_run(h)
    end = paras[-1].add_run()
    end._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(_el("w:updateFields", **{"w:val": "true"}))


def _cm_widths(pcts: list[int], minimum: float = 1.7) -> list[float]:
    """Column widths in cm. Word pads cells more than the PDF, so no column is narrower than a reference like F-001."""
    cm = [max(minimum, TEXT_WIDTH_CM * p / 100) for p in pcts]
    extra = sum(cm) - TEXT_WIDTH_CM
    wide = [i for i, w in enumerate(cm) if w > minimum]
    room = sum(cm[i] - minimum for i in wide) or 1
    return [w - extra * (w - minimum) / room if i in wide and extra > 0 else w for i, w in enumerate(cm)]


def _table(doc, rows: list[list[tuple[str, object]]], report: bool) -> None:
    cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    texts = [[inline.content for _, inline in r] for r in rows]
    widths = _cm_widths(column_widths(texts))
    for col, w in zip(tbl.columns, widths):
        col.width = Cm(w)
    header = texts[0]
    scales = [bar_scale(h) for h in header]
    for ri, r in enumerate(rows):
        for ci, (kind, inline) in enumerate(r):
            cell = tbl.cell(ri, ci)
            if ci < len(widths):
                cell.width = Cm(widths[ci])
            cell_p = cell.paragraphs[0]
            cell_p.paragraph_format.space_after = Pt(0)
            _inline(cell_p, inline)
            for run in cell_p.runs:
                run.font.size = Pt(9)
            if kind == "th_open":
                _shade(cell, HEADER_FILL)
                for run in cell_p.runs:
                    run.bold, run.font.color.rgb = True, NAVY
                continue
            level = from_labelled(inline.content)
            if level is not None:
                _shade(cell, level.fill)
                for run in cell_p.runs:
                    run.bold, run.font.color.rgb = True, RGBColor.from_string(level.ink)
            elif report and ci < len(scales) and scales[ci] and bar_pct(inline.content, scales[ci]) is not None:
                for run in cell_p.runs:
                    run.bold = True
    _repeat_header(tbl.rows[0])


def to_docx(md: str, title: str, meta: dict) -> bytes:
    doc = Document()
    _styles(doc)
    doc.core_properties.title = title
    doc.core_properties.author = "Futran Solutions"
    _running(doc, title, meta)
    report = bool(meta.get("cover"))
    if report and LOGO.is_file():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(60)
        p.add_run().add_picture(str(LOGO), height=Cm(1.4))

    toks = tokens(md)
    headings = [toks[i + 1].content for i, t in enumerate(toks) if t.type == "heading_open" and t.tag == "h2"]
    lists: list[str] = []
    table: list[list[tuple[str, object]]] | None = None
    row: list | None = None
    in_quote = False
    i = 0
    while i < len(toks):
        t = toks[i]
        if t.type == "heading_open":
            level = min(int(t.tag[1]), 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            _inline(p, toks[i + 1])
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("Bullet" if t.type == "bullet_list_open" else "Number")
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t.type == "blockquote_open":
            in_quote = True
        elif t.type == "blockquote_close":
            in_quote = False
        elif t.type == "table_open":
            table = []
        elif t.type == "tr_open":
            row = []
        elif t.type in ("th_open", "td_open"):
            row.append((t.type, toks[i + 1]))
            i += 3
            continue
        elif t.type == "tr_close":
            table.append(row)
        elif t.type == "table_close":
            _table(doc, table, report)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            table = None
        elif t.type == "paragraph_open":
            inline = toks[i + 1]
            if inline.content.strip() == SENTINEL:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif inline.content.strip() == TOC_SENTINEL:
                _toc(doc, headings)
            else:
                style = None
                if lists:
                    style = f"List {lists[-1]}" + (f" {min(len(lists), 3)}" if len(lists) > 1 else "")
                elif in_quote:
                    style = "Quote"
                p = doc.add_paragraph(style=style)
                _inline(p, inline)
            i += 3
            continue
        elif t.type == "hr":
            doc.add_paragraph()
        elif t.type in ("fence", "code_block"):
            p = doc.add_paragraph()
            r = p.add_run(t.content.rstrip("\n"))
            r.font.name = "Consolas"
        i += 1
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
